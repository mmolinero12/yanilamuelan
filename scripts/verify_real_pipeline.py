"""Prueba multipart, persistencia y Gemini reales con documentos pequeños.

Consume cuota de Gemini. Ejecutar con el Python del entorno backend.
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
import sys
import os
import socket
import subprocess
import tempfile
import time
from urllib.request import urlopen

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

from fastapi.testclient import TestClient
from backend.app import app

TEXT = (
    "Una VCN (Virtual Cloud Network) es una red privada virtual en Oracle Cloud Infrastructure. "
    "Una VCN contiene subredes. Las subredes alojan recursos como instancias de computo. "
    "Las reglas de seguridad controlan el trafico entrante y saliente. "
    "Las tablas de rutas determinan a donde se dirige el trafico de la red."
)


def sample(extension):
    buffer = BytesIO()
    if extension == "docx":
        from docx import Document
        document = Document()
        document.add_heading("Redes VCN", 0)
        document.add_paragraph(TEXT)
        document.save(buffer)
    elif extension == "pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Redes VCN"
        slide.placeholders[1].text = TEXT
        document.save(buffer)
    else:
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 700 Td ({TEXT}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(buffer)
    return buffer.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--formats", nargs="+", choices=["pdf", "docx", "pptx"],
                        default=["pdf", "docx", "pptx"])
    parser.add_argument("--streamlit", action="store_true",
                        help="Verifica Streamlit con un Uvicorn temporal y APIs reales.")
    args = parser.parse_args()
    if args.streamlit:
        return verify_streamlit()
    with TestClient(app) as client:
        for extension in args.formats:
            data = sample(extension)
            response = client.post("/api/v1/adapt", data={
                "perfil_destinatario": "Principiante", "formato_salida": "Quiz",
                "nicho_sector": "General", "nivel_detalle": "Didactico",
            }, files={"archivo": (f"verificacion-real.{extension}", data)})
            if response.status_code != 200:
                print(extension, "FALLO", response.status_code, response.json())
                return 1
            body = response.json()
            local = body["almacenamiento_local"]
            assert Path(local["documento_original"]).read_bytes() == data
            assert json.loads(Path(local["resultado_generado"]).read_text(encoding="utf-8")) == body
            assert body["metadatos"]["fragmentos_recuperados"] > 0
            assert body["evaluacion_calidad"]["claridad_pedagogica"] != "No evaluada"
            print(extension, "OK", "doc_id=" + body["metadatos"]["doc_id"],
                  "items=" + str(len(body["contenido_adaptado"]["items"])))
    return 0


def verify_streamlit():
    root = Path(__file__).resolve().parents[1]
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    environment = dict(os.environ, BACKEND_URL=f"http://127.0.0.1:{port}", PYTHONUTF8="1")
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    process = subprocess.Popen([
        sys.executable, "-X", "utf8", "-m", "uvicorn", "backend.app:app",
        "--host", "127.0.0.1", "--port", str(port),
    ], cwd=root, env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
       creationflags=flags)
    try:
        from backend.core.config import settings
        for _ in range(60):
            if process.poll() is not None:
                raise RuntimeError("El backend temporal no pudo arrancar.")
            try:
                with urlopen(f"http://127.0.0.1:{port}{settings.api_prefix}/health", timeout=1):
                    break
            except OSError:
                time.sleep(0.5)
        else:
            raise RuntimeError("El backend temporal no respondió a health.")
        frontend_python = root / ".venv-frontend" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory) / "verificacion-streamlit.docx"
            original.write_bytes(sample("docx"))
            resultado = subprocess.run([
                str(frontend_python), "-X", "utf8", str(root / "scripts/verify_streamlit_real.py"),
                str(original),
            ], cwd=root, env=environment, timeout=420, creationflags=flags)
            print("Streamlit HTTP real OK" if resultado.returncode == 0 else "Streamlit FALLO")
            return resultado.returncode
    finally:
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=10)


if __name__ == "__main__":
    raise SystemExit(main())
