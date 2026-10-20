"""Extrae texto del original persistido, sin alterar sus bytes."""
from pathlib import Path
from zipfile import ZipFile
from backend.core.exceptions import PipelineError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def _table_text(table) -> list[str]:
    return [" | ".join(cell.text for cell in row.cells) for row in table.rows]


def _shape_text(shapes) -> list[str]:
    parts = []
    for shape in shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.extend(_table_text(shape.table))
        if hasattr(shape, "shapes"):
            parts.extend(_shape_text(shape.shapes))
    return parts


def extraer_texto(path: Path) -> str:
    try:
        extension = path.suffix.lower()
        if extension in {".docx", ".pptx"}:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("Documento demasiado grande al descomprimir.")
        if extension == ".pdf":
            from pypdf import PdfReader
            texto = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif extension == ".docx":
            from docx import Document
            document = Document(path)
            texto = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [line for table in document.tables for line in _table_text(table)]
            )
        elif extension == ".pptx":
            from pptx import Presentation
            texto = "\n\n".join(
                "\n".join(_shape_text(slide.shapes)) for slide in Presentation(path).slides
            )
        elif extension in {".txt", ".md"}:
            texto = path.read_text(encoding="utf-8-sig")
        else:
            raise ValueError("Formato no admitido.")
    except Exception as error:
        raise PipelineError("extraccion", "No se pudo extraer el texto. Comprueba el formato, "
                            "la integridad y el tamaño del documento.", 422) from error
    if not texto.strip():
        raise PipelineError("extraccion", "El documento no contiene texto extraíble. "
                            "Los documentos escaneados requieren OCR, que no está incluido.", 422)
    return texto.strip()
