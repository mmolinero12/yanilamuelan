"""Prueba el HTTP y disco reales; solo sustituye llamadas externas a Gemini."""
from io import BytesIO
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from backend.app import app
from backend.core.config import settings
from shared.module_loader import load_module
from shared.contratos import ContenidoQuiz, ItemQuiz
from evaluacion.schemas.juez import ResultadoJuez, VeredictoItem
from evaluacion.services import evaluador

TEXT = "Una VCN es una red privada virtual de Oracle Cloud Infrastructure."
PARAMS = {"perfil_destinatario": "Principiante", "formato_salida": "Quiz Interactivo con Justificaciones",
          "nicho_sector": "General", "nivel_detalle": "Didactico"}


def document_bytes(extension):
    buffer = BytesIO()
    if extension == ".docx":
        from docx import Document
        document = Document()
        document.add_paragraph(TEXT)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Tabla VCN"
        document.save(buffer)
    elif extension == ".pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "VCN"
        slide.placeholders[1].text = TEXT
        document.save(buffer)
    elif extension == ".pdf":
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 700 Td ({TEXT}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(buffer)
    else:
        return TEXT.encode("utf-8")
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "documents_path", str(tmp_path / "originales"))
    monkeypatch.setattr(settings, "generated_path", str(tmp_path / "resultados"))
    return TestClient(app)


@pytest.fixture
def gemini_offline(tmp_path, monkeypatch):
    """El chunking, Chroma persistente, recuperación y evaluador siguen siendo reales."""
    vector_db = load_module("ingesta_rag", "models.vector_db")
    embeddings = load_module("ingesta_rag", "services.embedding_service")
    monkeypatch.setattr(vector_db, "CHROMA_PERSIST_DIR", str(tmp_path / "chroma"))

    class RemoteEmbeddings:
        def embed_documents(self, texts):
            return [[1.0, float("VCN" in text), 0.5] for text in texts]

        def embed_query(self, text):
            return [1.0, 1.0, 0.5]

    monkeypatch.setattr(embeddings.GeminiEmbeddingsResiliente, "_obtener_instancia_embeddings",
                        lambda self: RemoteEmbeddings())
    generator = load_module("orquestacion_agentes", "main")
    monkeypatch.setattr(generator, "_invocar_con_rotacion", lambda *args: ContenidoQuiz(
        titulo="Redes VCN", introduccion_contextualizada="Repaso de redes privadas.",
        items=[ItemQuiz(pregunta="¿Qué es una VCN?", opciones=["Red privada", "Disco"],
                        respuesta_correcta="Red privada", justificacion=TEXT)],
        conceptos_clave=["VCN"], tiempo_estimado_estudio_minutos=2,
    ))
    monkeypatch.setattr(evaluador, "_invocar_con_rotacion", lambda prompt: ResultadoJuez(
        veredictos_items=[VeredictoItem(indice=0, veredicto="respaldado", justificacion=TEXT)],
        claridad_pedagogica="Alta", observaciones="Información respaldada.",
    ))


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".pptx"])
def test_multipart_persiste_original_chroma_y_resultado(client, gemini_offline, extension):
    original = document_bytes(extension)
    response = client.post("/api/v1/adapt", data=PARAMS,
                           files={"archivo": ("vcn" + extension, original)})
    assert response.status_code == 200, response.json()
    body = response.json()
    local = body["almacenamiento_local"]
    assert Path(local["documento_original"]).read_bytes() == original
    assert Path(local["documento_original"]).suffix == extension
    assert json.loads(Path(local["resultado_generado"]).read_text(encoding="utf-8")) == body
    assert body["evaluacion_calidad"]["anclaje_fuente_score"] == 1.0
    assert body["metadatos"]["fragmentos_recuperados"] >= 1
    assert "almacenamiento_oci" not in body
    vector_db = load_module("ingesta_rag", "models.vector_db")
    indexing = load_module("ingesta_rag", "services.indexing_service")
    store = vector_db._obtener_vectorstore(indexing._embedding_service)
    records = store.get(where={"doc_id": body["metadatos"]["doc_id"]})
    assert TEXT in "\n".join(records["documents"])
    assert (Path(vector_db.CHROMA_PERSIST_DIR) / "chroma.sqlite3").is_file()


@pytest.mark.parametrize("name,data,status", [
    ("vcn.exe", b"data", 415), ("vcn.pdf", b"", 413),
    ("vcn.pdf", b"archivo corrupto", 422),
])
def test_upload_invalido_no_devuelve_exito(client, name, data, status):
    response = client.post("/api/v1/adapt", data=PARAMS, files={"archivo": (name, data)})
    assert response.status_code == status
    assert "contenido_adaptado" not in response.json()


def test_evaluacion_fallida_no_se_presenta_como_exito(client, gemini_offline, monkeypatch):
    from shared.contratos import EvaluacionCalidad
    from backend.services.evaluation_service import EvaluationService
    monkeypatch.setattr(EvaluationService, "evaluar", lambda *args: EvaluacionCalidad(
        anclaje_fuente_score=0, claridad_pedagogica="No evaluada", observaciones="Sin cuota"))
    response = client.post("/api/v1/adapt", data=PARAMS,
                           files={"archivo": ("vcn.txt", TEXT.encode())})
    assert response.status_code == 502
    assert response.json()["detail"]["etapa"] == "evaluacion"
    assert not Path(settings.generated_path).exists()


def test_error_externo_no_filtra_detalles(client, monkeypatch):
    from backend.services.rag_service import RagService
    def fail(*args):
        raise RuntimeError("credencial-que-no-debe-aparecer")
    monkeypatch.setattr(RagService, "indexar", fail)
    response = client.post("/api/v1/adapt", data=PARAMS,
                           files={"archivo": ("vcn.txt", TEXT.encode())})
    assert response.status_code == 502
    assert response.json()["detail"]["etapa"] == "indexacion"
    assert "credencial-que-no-debe-aparecer" not in response.text


def test_json_antiguo_se_rechaza(client):
    assert client.post("/api/v1/adapt", json=PARAMS).status_code == 422
